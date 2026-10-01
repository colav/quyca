import os
import zipfile
from typing import Any, IO, Iterator
from io import BytesIO

from openpyxl import load_workbook
from openpyxl.utils.exceptions import InvalidFileException

from quyca.domain.constants.project_file_schema import (
    ALLOWED_EXTENSIONS,
    PROJECT_FILE_COLUMNS,
    normalize_header,
)
from quyca.domain.exceptions.project_file_exceptions import ProjectFileException
from quyca.domain.models.project_file_model import ProjectFileContent, ProjectFileRow

_COLUMNS_BY_HEADER = {normalize_header(spec.name): spec.name for spec in PROJECT_FILE_COLUMNS}


def read_project_file(file_stream: BytesIO | IO[bytes], filename: str) -> ProjectFileContent:
    extension = os.path.splitext(filename or "")[1].lower()
    if extension not in ALLOWED_EXTENSIONS:
        raise ProjectFileException(
            f"Extensión '{extension or 'desconocida'}' no soportada. Formatos admitidos: {', '.join(ALLOWED_EXTENSIONS)}."
        )
    try:
        workbook = load_workbook(file_stream, read_only=True, data_only=True)
    except (InvalidFileException, zipfile.BadZipFile, KeyError, OSError) as error:
        raise ProjectFileException("El archivo no es un libro de Excel (.xlsx) válido.") from error
    try:
        if not workbook.worksheets:
            raise ProjectFileException("El libro de Excel no tiene hojas.")
        rows_iterator = workbook.worksheets[0].iter_rows(min_row=1, values_only=True)
        header_row = next(rows_iterator, None)
        if header_row is None or all(clean_value(cell) is None for cell in header_row):
            raise ProjectFileException("La primera fila de la hoja debe contener los encabezados.")
        index_to_column, ignored, duplicated = _map_headers(header_row)
        rows = read_rows(rows_iterator, index_to_column)
        return ProjectFileContent(
            sheet_count=len(workbook.worksheets),
            columns=list(index_to_column.values()),
            ignored_columns=ignored,
            duplicated_columns=duplicated,
            rows=rows,
        )
    finally:
        workbook.close()


def _map_headers(header_row: tuple) -> tuple[dict[int, str], list[str], list[str]]:
    index_to_column: dict[int, str] = {}
    ignored: list[str] = []
    duplicated: list[str] = []
    for index, cell in enumerate(header_row):
        header = clean_value(cell)
        if header is None:
            continue
        canonical = _COLUMNS_BY_HEADER.get(normalize_header(str(header)))
        if canonical is None:
            ignored.append(str(header))
        elif canonical in index_to_column.values():
            duplicated.append(canonical)
        else:
            index_to_column[index] = canonical
    return index_to_column, ignored, duplicated


def read_rows(rows_iterator: Iterator, index_to_column: dict[int, str]) -> list[ProjectFileRow]:
    rows = []
    for row_number, cells in enumerate(rows_iterator, start=2):
        values = {
            column: clean_value(cells[index] if index < len(cells) else None)
            for index, column in index_to_column.items()
        }
        if all(value is None for value in values.values()):
            continue
        rows.append(ProjectFileRow(row_number=row_number, values=values))
    return rows


def clean_value(value: Any) -> Any:
    if isinstance(value, str):
        value = value.strip()
        return value or None
    return value
